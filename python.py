from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

slides_data = [
    ("CurryKrust: Curriculum + Trust", [
        "The Mastery-Gated AI Project Accelerator",
        "Inputs: Project Name, Description, Daily Time, Deadline, Current Skill Level",
        "Stack: HTML5/CSS3/JS Frontend + FastAPI + Tavily API + Scikit-Learn ML + Gemini AI"
    ]),
    ("1. The Problem: Why Self-Learners Miss Deadlines", [
        "Tutorial Hell: Search engines dump thousands of unverified links with zero pacing.",
        "AI Hallucinations: Standard LLMs invent fake links and let users passively skim.",
        "No Mastery Check: Learners jump ahead without proving they can build the basics."
    ]),
    ("2. The CurryKrust Module Gauntlet", [
        "Every generated module enforces a strict 4-step mastery loop:",
        "1. Credible Theory: Domain-verified docs & videos ranked by Machine Learning.",
        "2. Immediate Practice: 5-minute hands-on micro-exercise right after reading.",
        "3. Gatekeeper Quiz + Mini-Project: 5 MCQs + a real piece of the user's project.",
        "4. Final AI Assessment: Must score >= 80% to unlock the next module in the UI."
    ]),
    ("3. Under the Hood: APIs + ML + AI + HTML/CSS/JS", [
        "External APIs (Tavily & YouTube): Pulls real tutorials from MDN, Python Docs, MIT, etc.",
        "Machine Learning (Scikit-Learn): TF-IDF & Cosine Similarity score resource relevance (0-100%).",
        "Generative AI (Gemini 2.5 Flash): Structures timed modules & grades project submissions.",
        "Frontend (HTML5, CSS3, Vanilla JS): Responsive dashboard with DOM lock/unlock states."
    ]),
    ("4. End-to-End Execution Roadmap", [
        "Phase 1: FastAPI Backend + Pydantic JSON Schemas + Deadline Hour Calculator.",
        "Phase 2: Tavily API Integration + Scikit-Learn Cosine Similarity Credibility Engine.",
        "Phase 3: Gemini Structured Curriculum Generator + Automated Rubric Evaluator.",
        "Phase 4: Custom HTML/CSS/JS Interface with Locked Module Cards & Live Telemetry."
    ])
]

for title_text, bullets in slides_data:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    # Dark background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(15, 23, 42)
    
    # Title styling
    title = slide.shapes.title
    title.text = title_text
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(234, 88, 12)
    title.text_frame.paragraphs[0].font.size = Pt(36)
    title.text_frame.paragraphs[0].font.bold = True

    # Body bullets
    tf = slide.placeholders[1].text_frame
    tf.clear()
    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(22)
        p.font.color.rgb = RGBColor(248, 250, 252)
        p.space_after = Pt(14)

prs.save("CurryKrust_Pitch_Deck.pptx")
print("Saved CurryKrust_Pitch_Deck.pptx successfully!")